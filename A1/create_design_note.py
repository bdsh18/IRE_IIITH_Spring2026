from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("outputs/design_note_assignment1.docx")

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr(); fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), color); props.append(fill)

def set_cell_text(cell, text, bold=False):
    cell.text = ""; p = cell.paragraphs[0]; r = p.add_run(text); r.bold = bold; r.font.size = Pt(9); p.paragraph_format.space_after = Pt(0); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, headers): shade(c, "E8EEF5"); set_cell_text(c, h, True)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row): set_cell_text(c, str(v))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.add_run(text); p.paragraph_format.space_after = Pt(2)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.add_run(text); return p

def main():
    doc = Document(); section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75); section.left_margin = section.right_margin = Inches(.8)
    section.header_distance = section.footer_distance = Inches(.35)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10); normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.05
    for level, size in [(1, 14), (2, 11)]:
        style = doc.styles[f"Heading {level}"]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(46,116,181); style.font.bold = True; style.paragraph_format.space_before = Pt(8); style.paragraph_format.space_after = Pt(3)
    header = section.header.paragraphs[0]; header.text = "CS4.406 - Information Retrieval & Extraction | Assignment 1"; header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor(100,100,100)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run("News Retrieval Design Note | "); field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("Lexical and Semantic Retrieval for News Recommendation"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor(11,37,69)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("Assignment 1 Design Note | CS4.406 Information Retrieval & Extraction"); r.italic = True; r.font.size = Pt(10)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Student: [Replace with your name] | Date: 23 August 2026").font.size = Pt(9)
    heading(doc, "1. Objective and System Overview")
    doc.add_paragraph("The task is to rank the articles displayed in a news impression by click likelihood. I built a reproducible pipeline for MIND (English) and EB-NeRD (Danish), using article text, a user’s previous clicks, and lightweight behavioural features. The project deliberately separates data preparation, candidate retrieval, offline evaluation, and leaderboard prediction generation.")
    table(doc, ["Component", "Implementation"], [["Development data", "MIND-small and EB-NeRD demo; large archives reserved for Codabench prediction."], ["Feature store", "Parquet article and chronological impression tables with click history and recency weights."], ["Lexical retrieval", "BM25 inverted index over title + abstract; query from five recent clicked titles."], ["Semantic retrieval", "MIND entity-vector mean pooling; EB-NeRD supplied Word2Vec document vectors; cosine nearest-neighbour search."], ["Leaderboard output", "Validated ZIP files with the required prediction.txt / predictions.txt only."]])
    heading(doc, "2. Reproducible Data Pipeline")
    doc.add_paragraph("A single `make data` command extracts the small development archives and writes unified Parquet feature stores. Each article has ID, title, abstract, body, category, entities, and embedding metadata. Each impression has timestamp, user ID, session ID, shown candidates, clicked items, click history, and exponentially decayed history weights.")
    bullet(doc, "Temporal split only: oldest impressions for train, then validation, then the newest data for test. No random interaction split is used.")
    bullet(doc, "Leakage protection: the pipeline represents only prior user history at recommendation time; future clicks are excluded.")

    doc.add_page_break(); heading(doc, "3. Retrieval Methods and Design Choices")
    heading(doc, "3.1 BM25 lexical baseline", 2)
    doc.add_paragraph("BM25 indexes title and abstract terms using an inverted index. The user query concatenates titles from the five most recent clicked articles. The index retrieves top-K articles for K = 50, 100, and 200. Exact keyword overlap is interpretable, inexpensive, and a strong first baseline for breaking news.")
    heading(doc, "3.2 Semantic retrieval", 2)
    doc.add_paragraph("For MIND, article vectors are created by averaging the available 100-dimensional entity embeddings from title and abstract annotations. For EB-NeRD, I load the provided 300-dimensional Word2Vec document vectors. The user vector is the mean of vectors for clicked history articles. Cosine similarity retrieves the nearest documents; brute-force search is used only for small/demo development data.")
    heading(doc, "3.3 Alternatives considered", 2)
    table(doc, ["Alternative", "Decision and rationale"], [["TF-IDF", "Not used separately because BM25 gives term-frequency saturation and length normalisation."], ["BERT/XLM-R embeddings", "Appropriate future upgrade, but supplied embeddings enabled reproducible CPU-scale experiments."], ["FAISS ANN", "Required at large scale; brute-force cosine is simpler and correct for the small/demo evaluation."], ["Global popularity only", "Used as a valid initial submission baseline, then improved with recency-weighted category/subcategory preferences."]])
    heading(doc, "4. Experiment Results")
    doc.add_paragraph("The following candidate-generation figures are from an equal 500-impression validation smoke test. They measure whether the ground-truth clicked article appears in the retrieved list.")
    table(doc, ["Dataset", "Method", "R@50", "R@100", "R@200"], [["MIND-small", "BM25", "1.10%", "2.44%", "4.16%"], ["MIND-small", "Semantic", "0.90%", "1.56%", "2.02%"], ["EB-NeRD demo", "BM25", "1.60%", "2.40%", "3.60%"], ["EB-NeRD demo", "Semantic", "4.00%", "5.00%", "6.40%"]])

    doc.add_page_break(); heading(doc, "5. Offline Evaluation")
    doc.add_paragraph("The evaluation harness ranks the candidates already displayed in each labelled validation impression. It reports AUC, MRR, nDCG@5, nDCG@10, intra-list diversity@10, novelty@10, coverage@10, cold/warm user slices, and 95% bootstrap confidence intervals.")
    table(doc, ["Dataset", "Method", "AUC (95% CI)", "MRR", "nDCG@10"], [["MIND", "BM25", "0.549 (0.523-0.573)", "0.255", "0.280"], ["MIND", "Semantic", "0.570 (0.545-0.594)", "0.255", "0.291"], ["EB-NeRD", "BM25", "0.495 (0.467-0.520)", "0.307", "0.427"], ["EB-NeRD", "Semantic", "0.490 (0.463-0.518)", "0.310", "0.424"]])
    heading(doc, "5.1 Observations", 2)
    bullet(doc, "MIND: semantic ranking slightly improved AUC and nDCG@10, although BM25 had stronger candidate-generation recall. Entity coverage limits MIND semantic vectors.")
    bullet(doc, "EB-NeRD: supplied document embeddings improved retrieval recall, but BM25 was marginally stronger in candidate-ranking AUC/nDCG in the sampled validation run. The confidence intervals overlap.")
    bullet(doc, "Warm MIND users outperformed cold users, confirming that longer click histories make user representations more reliable.")
    heading(doc, "5.2 Codabench submissions", 2)
    doc.add_paragraph("Both large-test submissions follow the prescribed rank-permutation format and preserve source row order. MIND used an improved personalised score combining decayed history category/subcategory affinity with global click popularity. Insert the two final Codabench leaderboard screenshots below before Moodle submission.")
    table(doc, ["Required evidence", "Insert before submission"], [["MIND leaderboard", "[Paste your MIND Codabench result screenshot here]"], ["EB-NeRD leaderboard", "[Paste your EB-NeRD Codabench result screenshot here]"]])

    doc.add_page_break(); heading(doc, "6. Scalability and 10x Limits")
    doc.add_paragraph("The small/demo pipeline is intentionally simple and reproducible. At 10x scale, several components become bottlenecks:")
    bullet(doc, "Brute-force cosine similarity becomes too slow and memory-intensive; replace it with FAISS/HNSW ANN indexes and batch queries.")
    bullet(doc, "Python row loops over tens of millions of impressions become expensive; use Polars/PyArrow scans, row groups, and streaming writers.")
    bullet(doc, "Full article-token inverted indexes and dense embeddings need disk-backed caching, versioned artifacts, and incremental refreshes for newly published news.")
    bullet(doc, "A single popularity score is too coarse. A production ranker should combine lexical, semantic, recency, category affinity, freshness, and candidate-level features in a learned model.")
    heading(doc, "7. Conclusion")
    doc.add_paragraph("This work establishes a measured retrieval foundation on two distinct news datasets. The pipeline is reproducible, avoids temporal leakage, supports lexical and semantic retrieval, and produces validated leaderboard files. The main lesson is that the best retrieval signal depends on dataset representation: exact text overlap remains useful, while document embeddings can capture semantic similarity when high-quality vectors are available.")
    heading(doc, "References", 2)
    doc.add_paragraph("Wu et al. (2020). MIND: A Large-scale Dataset for News Recommendation. ACL.\nKruse et al. (2024). EB-NeRD: A Large-scale Dataset for News Recommendation. RecSys Challenge.")
    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)

if __name__ == "__main__": main()
